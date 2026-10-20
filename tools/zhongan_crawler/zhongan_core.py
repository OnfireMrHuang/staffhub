"""Response validation, cursor pagination and atomic Excel upserts."""
from __future__ import annotations

import hashlib
import json
import os
import tempfile
from pathlib import Path

from filelock import FileLock
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class CollectionError(RuntimeError):
    pass


class LoginExpired(CollectionError):
    pass


def response_data(status, payload):
    if status == 401:
        raise LoginExpired("登录已过期，请重新登录。")
    if not 200 <= status < 300:
        raise CollectionError(f"接口 HTTP {status}，本轮不更新 Excel。")
    if not isinstance(payload, dict):
        raise CollectionError("接口响应必须是 JSON 对象。")
    codes = {str(payload.get(k)) for k in ("code", "returnCode")}
    if codes & {"401", "1013", "1014"}:
        raise LoginExpired("众安会话已过期，请重新登录。")
    if payload.get("success") is False or not (payload.get("success") is True or "200" in codes):
        raise CollectionError("接口业务失败，响应未被标记为成功。")
    if "data" not in payload or payload["data"] is None:
        raise CollectionError("接口未返回 data，不能作为空数据导出。")
    return payload["data"]


def collect_orders(fetch, base, max_pages=10000):
    """getOrderList returns a list; the LAST record carries the next cursor."""
    result, seen_orders, seen_cursors = [], set(), set()
    cursor = None
    for _ in range(max_pages):
        payload = dict(base)
        if cursor is not None:
            payload["nextPage"] = cursor
        rows = fetch(payload)
        if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows):
            raise CollectionError("订单 data 格式变化：预期列表。")
        if not rows:
            return result
        for row in rows:
            identity = row.get("orderNo")
            if identity is None or not str(identity).strip():
                raise CollectionError("订单缺少 orderNo，无法可靠增量更新。")
            identity = str(identity)
            if identity in seen_orders:
                raise CollectionError("分页出现重复订单，请重新采集以避免漏单。")
            seen_orders.add(identity)
        result.extend(rows)
        if "nextPage" not in rows[-1]:
            raise CollectionError("末条订单缺少 nextPage，无法确认已采集完整。")
        cursor = rows[-1]["nextPage"]
        if cursor in (None, "", False, 0):
            return result
        if isinstance(cursor, bool) or not isinstance(cursor, (str, int)):
            raise CollectionError("nextPage 游标类型变化。")
        marker = str(cursor)
        if marker in seen_cursors:
            raise CollectionError("接口返回重复分页游标。")
        seen_cursors.add(marker)
    raise CollectionError(f"订单超过 {max_pages} 页，本轮不更新 Excel。")


def flatten(value, prefix=""):
    result = {}
    for key, item in value.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(item, dict) and item:
            result.update(flatten(item, name))
        else:
            result[name] = item
    return result


def record_key(*parts):
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False).encode()).hexdigest()


def cell_value(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, int) and abs(value) >= 10**15:
        value = str(value)
    if isinstance(value, str):
        if len(value) > 32767:
            raise CollectionError("某个字段超出 Excel 单元格上限，未写入，完整数据见原始 JSON。")
        if any(ord(c) < 32 and c not in "\t\n\r" for c in value):
            raise CollectionError("接口字段包含 Excel 不支持的控制字符。")
    return value


def upsert_sheet(workbook, name, records):
    ws = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
    headers = [c.value for c in ws[1]] if ws.cell(1, 1).value is not None else []
    for record in records:
        for field in record:
            if field not in headers:
                headers.append(field)
    # Keep an empty order sheet usable, including when the first collection has no orders.
    if not headers:
        headers = ["记录键", "直播平台", "直播间号码", "统计日期", "直播场次ID", "订单号",
                   "订单人", "金额（接口原值）", "订单时间", "状态码", "订单状态",
                   "首次采集时间", "最近采集时间"]
    for col, title in enumerate(headers, 1):
        ws.cell(1, col, title)
    if headers[0] != "记录键":
        raise CollectionError(f"{name} 工作表结构不兼容，记录键必须在第一列。")
    existing = {}
    for row in range(2, ws.max_row + 1):
        key = ws.cell(row, 1).value
        if not key or key in existing:
            raise CollectionError(f"{name} 存在空键或重复键，请先修复工作簿。")
        existing[key] = row
    first_col = headers.index("首次采集时间") + 1
    for record in records:
        key = record["记录键"]
        row = existing.get(key)
        if row is None:
            row = ws.max_row + 1
            existing[key] = row
        else:
            record = {**record, "首次采集时间": ws.cell(row, first_col).value}
        for col, field in enumerate(headers, 1):
            # Preserve columns that this API no longer returns.
            if field in record:
                value = cell_value(record[field])
                cell = ws.cell(row, col)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = "s"  # Treat API strings as text, never Excel formulas.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="224B70")
        ws.column_dimensions[get_column_letter(cell.column)].width = min(42, max(18, len(str(cell.value)) + 6))
    ws.column_dimensions["A"].hidden = True


def export_excel(path: Path, batches, collected_at):
    dashboards, orders = [], []
    for room, aggregate, rows in batches:
        platform, dy_no, day, room_id = (str(room[k]) for k in ("livePlatform", "dyNo", "startDate", "roomId"))
        common = {"直播平台": platform, "直播间号码": dy_no, "统计日期": day, "直播场次ID": room_id,
                  **flatten(room, "直播间"), "首次采集时间": collected_at, "最近采集时间": collected_at}
        dashboards.append({"记录键": record_key(platform, dy_no, day, room_id), **common,
                           **flatten(aggregate, "看板")})
        for order in rows:
            if not order.get("orderNo"):
                raise CollectionError("订单缺少 orderNo。")
            holder = order.get("holderInfo") or {}
            status = order.get("orderStatus")
            orders.append({"记录键": record_key(platform, dy_no, room_id, str(order["orderNo"])), **common,
                           "订单号": str(order["orderNo"]),
                           "订单人": holder.get("name") if isinstance(holder, dict) else None,
                           "金额（接口原值）": order.get("orderAmount"),
                           "订单时间": " ".join(str(order[k]) for k in ("bizDate", "bizTime") if order.get(k)),
                           "状态码": status,
                           "订单状态": {"2": "已支付", "3": "已支付", "-1": "已退保"}.get(str(status), f"未支付/其他({status})"),
                           **flatten(order, "订单")})
    path.parent.mkdir(parents=True, exist_ok=True)
    with FileLock(str(path) + ".lock", timeout=10):
        wb = load_workbook(path) if path.exists() else Workbook()
        temp = None
        try:
            if not path.exists():
                wb.remove(wb.active)
            upsert_sheet(wb, "看板统计", dashboards)
            upsert_sheet(wb, "订单明细", orders)
            fd, temp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
            os.close(fd)
            wb.save(temp)
            os.replace(temp, path)
        finally:
            wb.close()
            if temp and os.path.exists(temp):
                os.unlink(temp)
