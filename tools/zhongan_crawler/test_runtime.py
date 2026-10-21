"""Local HTTP integration tests; never query Zhongan or send a real SMS."""
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

import zhongan_collect as crawler
from zhongan_core import CollectionError, LoginExpired


class LocalHandler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_POST(self):
        payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        self.server.requests.append((self.path, payload, dict(self.headers)))
        status, body = self.server.respond(self.path, payload, self.headers)
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(body).encode())


class RuntimeTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), LocalHandler)
        self.server.requests = []
        self.server.respond = lambda *_: (200, {"success": True})
        thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)
        self.origin = f"http://127.0.0.1:{self.server.server_port}"

    def test_sms_success_is_throttled_across_calls(self):
        state = self.root / "sms.json"
        self.assertTrue(crawler.notify_expired("13800000000", state, gateway=self.origin, token="test"))
        self.assertFalse(crawler.notify_expired("13800000000", state, gateway=self.origin, token="test"))
        self.assertEqual(len(self.server.requests), 1)
        _, payload, headers = self.server.requests[0]
        self.assertEqual(payload["phone"], "13800000000")
        self.assertEqual(payload["event"], "zhongan_login_expired")
        self.assertEqual(headers["Authorization"], "Bearer test")

    def test_sms_failure_does_not_start_cooldown(self):
        self.server.respond = lambda *_: (200, {"success": False})
        state = self.root / "sms.json"
        with self.assertRaises(CollectionError):
            crawler.notify_expired("13800000000", state, gateway=self.origin)
        self.assertFalse(state.exists())

    def test_sms_configuration_error_does_not_replace_login_failure(self):
        args = SimpleNamespace(phone="13800000000", session=self.root / "session.json")
        with patch.dict("os.environ", {"ZHONGAN_SMS_GATEWAY_URL": self.origin,
                                       "ZHONGAN_SMS_COOLDOWN_SECONDS": "invalid"}):
            crawler.try_notify(args)
        self.assertEqual(self.server.requests, [])

    def test_full_collection_reuses_auth_and_updates_same_excel_rows(self):
        self.paid = False

        def respond(path, payload, headers):
            if "sid=test-cookie" not in headers.get("Cookie", "") or headers.get("X-Usercenter-Session") != "test-session":
                return 200, {"code": "1013"}
            name = path.rsplit("/", 1)[-1]
            if name == "getLiveDy":
                data = [{"dyNo": "37365205524", "name": "Synthetic room"}]
            elif name == "getLiveAttribute":
                data = {"roomId": "9001", "attributeOfflineDTOList": []}
            elif name == "getDashboardAggregate":
                if payload != {"livePlatform": "DOUYIN", "dyNo": "37365205524", "roomIdList": ["9001"]}:
                    return 400, {}
                data = {"orderCnt": 2, "totalPremium": 30 if self.paid else 10}
            elif name == "getOrderList":
                if payload.get("showDetail") != "Y" or payload.get("startDate") != "2026-10-20":
                    return 400, {}
                data = [{"orderNo": "002" if payload.get("nextPage") == 37 else "001",
                         "holderInfo": {"name": "Synthetic user"}, "orderAmount": 10,
                         "orderStatus": 2 if self.paid else 1,
                         "nextPage": None if payload.get("nextPage") == 37 else 37}]
            else:
                return 404, {}
            return 200, {"success": True, "data": data}

        self.server.respond = respond
        state = {"cookies": [{"name": "sid", "value": "test-cookie", "domain": "127.0.0.1",
                              "path": "/", "expires": -1, "httpOnly": True, "secure": False, "sameSite": "Lax"}],
                 "origins": [{"origin": self.origin, "localStorage": [
                     {"name": "ATLANTIS_SESSION_ID", "value": "test-session"},
                     {"name": "SERVICE_NAME", "value": "test-service"}]}]}
        args = SimpleNamespace(output=self.root / "report.xlsx", dy_no="37365205524",
                               date="2026-10-20", platform="DOUYIN", all_sessions=False,
                               room_id=None, max_pages=10)
        with patch.object(crawler, "ORIGIN", self.origin), sync_playwright() as pw:
            api = pw.request.new_context(storage_state=state)
            try:
                context = SimpleNamespace(request=api, storage_state=api.storage_state)
                client = crawler.Client(context, self.origin, pause=0)
                crawler.collect_once(client, args)
                self.paid = True
                crawler.collect_once(client, args)
                wb = load_workbook(args.output)
                try:
                    self.assertEqual(wb["看板统计"].max_row, 2)
                    self.assertEqual(wb["订单明细"].max_row, 3)
                    headers = [c.value for c in wb["订单明细"][1]]
                    status_col = headers.index("状态码") + 1
                    self.assertEqual(wb["订单明细"].cell(2, status_col).value, 2)
                finally:
                    wb.close()
                before = args.output.read_bytes()
                self.server.respond = lambda *_: (200, {"returnCode": "1014"})
                with self.assertRaises(LoginExpired):
                    crawler.collect_once(client, args)
                self.assertEqual(args.output.read_bytes(), before)
                manifests = [json.loads(p.read_text()) for p in self.root.glob("raw/*/manifest.json")]
                self.assertEqual(sum(m["complete"] for m in manifests), 2)
                self.assertEqual(sum(not m["complete"] for m in manifests), 1)
            finally:
                api.dispose()


class RawSnapshotTests(unittest.TestCase):
    def test_raw_snapshot_is_complete_without_writing_excel(self):
        class FakeClient:
            raw_dir = None

            def post(self, endpoint, payload):
                return {"rooms": [{"dyNo": "001"}],
                        "attribute": {"roomId": "9001"},
                        "aggregate": {"orderCnt": 0}, "orders": []}[endpoint]

        with tempfile.TemporaryDirectory() as directory:
            args = SimpleNamespace(output=Path(directory) / "unused.xlsx", dy_no="001",
                                   date="2026-10-21", platform="DOUYIN", room_id=None,
                                   all_sessions=False, max_pages=10, raw_only=True)
            client = FakeClient()
            with patch.object(crawler, "export_excel") as exporter:
                crawler.collect_once(client, args)
                exporter.assert_not_called()
            run = next((Path(directory) / "raw").iterdir())
            manifest = json.loads((run / "manifest.json").read_text())
            snapshot = json.loads((run / "batches.json").read_text())
            self.assertTrue(manifest["complete"])
            self.assertIsNone(manifest["workbook"])
            self.assertEqual(snapshot["batches"][0][0]["dyNo"], "001")
            self.assertEqual(snapshot["batches"][0][2], [])
            self.assertFalse(args.output.exists())
            self.assertIsNone(client.raw_dir)

    def test_raw_snapshot_failure_is_not_marked_complete(self):
        class FailedClient:
            raw_dir = None

            def post(self, endpoint, payload):
                raise LoginExpired("expired")

        with tempfile.TemporaryDirectory() as directory:
            args = SimpleNamespace(output=Path(directory) / "unused.xlsx", dy_no="001",
                                   date="2026-10-21", platform="DOUYIN", raw_only=True)
            client = FailedClient()
            with self.assertRaises(LoginExpired):
                crawler.collect_once(client, args)
            run = next((Path(directory) / "raw").iterdir())
            self.assertFalse(json.loads((run / "manifest.json").read_text())["complete"])
            self.assertFalse((run / "batches.json").exists())
            self.assertIsNone(client.raw_dir)


class SessionSelectionTests(unittest.TestCase):
    def test_default_current_all_and_explicit_historical_session(self):
        attr = {"roomId": "current", "attributeOfflineDTOList": [{"roomId": "old"}]}
        args = SimpleNamespace(room_id=None, all_sessions=False)
        self.assertEqual([s["roomId"] for s in crawler.select_sessions(attr, args)], ["current"])
        args.all_sessions = True
        self.assertEqual([s["roomId"] for s in crawler.select_sessions(attr, args)], ["current", "old"])
        args.room_id = ["old"]
        self.assertEqual([s["roomId"] for s in crawler.select_sessions(attr, args)], ["old"])
        args.room_id = ["missing"]
        with self.assertRaises(CollectionError):
            crawler.select_sessions(attr, args)

    def test_empty_sessions_never_select_another_room(self):
        with self.assertRaises(CollectionError):
            crawler.select_sessions({}, SimpleNamespace(room_id=None, all_sessions=False))


class ImportSessionTests(unittest.TestCase):
    def test_chrome_cookie_header_and_token_import(self):
        url = crawler.ORIGIN + crawler.API_ROOT + "getDashboardAggregate"
        state = crawler.session_from_curl(
            f"curl '{url}' \\\n -b 'sid=abc=def; second=xyz' "
            "-H 'x-usercenter-session: session-token' -H 'x-service-name: undefined'")
        self.assertEqual(len(state["cookies"]), 2)
        self.assertEqual(state["cookies"][0]["value"], "abc=def")
        self.assertEqual(state["cookies"][0]["domain"], "adagent.zhongan.com")
        self.assertEqual(state["origins"][0]["localStorage"],
                         [{"name": "ATLANTIS_SESSION_ID", "value": "session-token"}])

    def test_sanitized_request_and_other_origins_rejected(self):
        url = crawler.ORIGIN + crawler.API_ROOT + "getOrderList"
        for text in (f"curl '{url}'", "curl https://example.com -b 'sid=abc'",
                     f"curl '{url}' https://example.com -b 'sid=abc'"):
            with self.subTest(text=text), self.assertRaises(CollectionError):
                crawler.session_from_curl(text)

    def test_shell_metacharacters_are_only_data(self):
        url = crawler.ORIGIN + crawler.API_ROOT + "getOrderList"
        state = crawler.session_from_curl(f"curl '{url}' -H 'Cookie: sid=$(do-not-run)'")
        self.assertEqual(state["cookies"][0]["value"], "$(do-not-run)")


if __name__ == "__main__":
    unittest.main()
