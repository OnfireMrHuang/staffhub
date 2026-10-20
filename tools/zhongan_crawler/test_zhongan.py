"""Synthetic fixtures based on the frontend contract; no customer data."""
import tempfile
import unittest
from pathlib import Path

import zhongan_core as core
from openpyxl import load_workbook


class ResponseTests(unittest.TestCase):
    def test_expired_business_session_is_not_an_empty_report(self):
        for payload in ({"returnCode": "401"}, {"code": "1013"}, {"code": "1014"}):
            with self.subTest(payload=payload), self.assertRaises(core.LoginExpired):
                core.response_data(200, payload)

    def test_forbidden_and_malformed_data_fail(self):
        with self.assertRaises(core.CollectionError):
            core.response_data(403, {})
        with self.assertRaises(core.CollectionError):
            core.response_data(200, {"success": False, "data": []})
        with self.assertRaises(core.CollectionError):
            core.response_data(200, {"success": True})

    def test_success_zero_and_empty_list_are_preserved(self):
        self.assertEqual(core.response_data(200, {"success": True, "data": []}), [])
        self.assertEqual(core.response_data(200, {"code": "200", "data": {"amount": 0}}), {"amount": 0})


class PaginationTests(unittest.TestCase):
    def test_follows_last_row_cursor_and_preserves_base_filter(self):
        calls = []

        def fetch(payload):
            calls.append(payload)
            if "nextPage" not in payload:
                return [{"orderNo": "0001", "nextPage": 73}]
            self.assertEqual(payload, {"dyNo": "37365205524", "showDetail": "Y", "nextPage": 73})
            return [{"orderNo": "0002", "nextPage": None}]

        rows = core.collect_orders(fetch, {"dyNo": "37365205524", "showDetail": "Y"})
        self.assertEqual([r["orderNo"] for r in rows], ["0001", "0002"])
        self.assertEqual(len(calls), 2)

    def test_cursor_loop_and_missing_order_identity_are_rejected(self):
        with self.assertRaises(core.CollectionError):
            core.collect_orders(lambda _: [{"orderNo": "1", "nextPage": 1}], {})
        with self.assertRaises(core.CollectionError):
            core.collect_orders(lambda _: [{"holderInfo": {"name": "Example"}}], {})

    def test_missing_cursor_is_not_silently_treated_as_last_page(self):
        with self.assertRaises(core.CollectionError):
            core.collect_orders(lambda _: [{"orderNo": "1"}], {})

    def test_page_limit_and_changed_duplicate_order_fail(self):
        with self.assertRaises(core.CollectionError):
            core.collect_orders(lambda _: [{"orderNo": "1", "nextPage": 2}], {}, max_pages=1)
        pages = iter([[{"orderNo": "1", "nextPage": 2}], [{"orderNo": "1", "nextPage": None}]])
        with self.assertRaises(core.CollectionError):
            core.collect_orders(lambda _: next(pages), {})

    def test_empty_orders_are_valid(self):
        self.assertEqual(core.collect_orders(lambda _: [], {}), [])


class ExcelTests(unittest.TestCase):
    def test_upsert_preserves_other_rooms_and_long_identifiers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            room = {"livePlatform": "DOUYIN", "dyNo": "37365205524", "startDate": "2026-10-20", "roomId": "9001"}
            one = {"orderNo": "001234567890123456789", "holderInfo": {"name": "=1+1"},
                   "orderAmount": 20.5, "orderStatus": 1, "bizTime": "2026-10-20 10:00:00"}
            core.export_excel(path, [(room, {"orderCnt": 1}, [one])], "first")
            other = {**room, "roomId": "9002"}
            core.export_excel(path, [(other, {"orderCnt": 0}, [])], "second")
            core.export_excel(path, [(room, {"orderCnt": 1, "newMetric": 0},
                                      [{**one, "orderStatus": 2}])], "third")
            wb = load_workbook(path)
            self.addCleanup(wb.close)
            self.assertEqual(wb["看板统计"].max_row, 3)
            ws = wb["订单明细"]
            self.assertEqual(ws.max_row, 2)
            values = dict(zip((c.value for c in ws[1]), (c.value for c in ws[2])))
            self.assertEqual(values["订单号"], "001234567890123456789")
            self.assertEqual(values["状态码"], 2)
            self.assertEqual(values["订单状态"], "已支付")
            self.assertEqual(values["首次采集时间"], "first")
            self.assertEqual(values["最近采集时间"], "third")
            self.assertEqual(values["金额（接口原值）"], 20.5)
            name_cell = ws.cell(2, [c.value for c in ws[1]].index("订单人") + 1)
            self.assertEqual(name_cell.data_type, "s")

    def test_invalid_batch_does_not_damage_existing_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            room = {"livePlatform": "DOUYIN", "dyNo": "1", "startDate": "2026-10-20", "roomId": "9"}
            core.export_excel(path, [(room, {"orderCnt": 0}, [])], "first")
            before = path.read_bytes()
            with self.assertRaises(core.CollectionError):
                core.export_excel(path, [(room, {}, [{"orderNo": "1", "bad": "x" * 32768}])], "second")
            self.assertEqual(path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
